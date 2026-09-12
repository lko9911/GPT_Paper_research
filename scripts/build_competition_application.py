from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "경진대회 참가신청서.docx"
OUTPUT = ROOT / "경진대회 참가신청서_작성본.docx"

FONT = "Malgun Gothic"
INK = "172033"
ACCENT = "1B5E70"
ACCENT_LIGHT = "E9F3F5"
MUTED = "5C667A"
LINE = "C9D6DA"
WARNING = "FFF2CC"
TABLE_WIDTH_DXA = 9440


def set_run_font(run, size: float = 10.5, bold: bool | None = None, color: str = INK) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 120, start: int = 140, bottom: int = 120, end: int = 140) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LINE, size: int = 6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError("Table widths must add up to the usable table width")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "140")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.2

    for name, size, color, before, after in (
        ("Heading 1", 12.5, ACCENT, 12, 7),
        ("Heading 2", 11.0, INK, 8, 4),
        ("Heading 3", 10.5, ACCENT, 6, 3),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(10.3)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.left_indent = Cm(0.75)
        style.paragraph_format.first_line_indent = Cm(-0.3)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.15


def add_section_heading(doc: Document, text: str, page_break: bool = False) -> None:
    if page_break:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.add_run(f"□ {text}")
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), ACCENT_LIGHT)
    p_pr.append(shd)
    spacing = p_pr.get_or_add_spacing()
    spacing.set(qn("w:before"), "120")
    spacing.set(qn("w:after"), "100")


def add_item_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Heading 2")
    paragraph.add_run(f"○ {text}")


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.45)
    paragraph.paragraph_format.keep_together = True
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True, color=ACCENT)
        run = paragraph.add_run(text[len(bold_lead):])
        set_run_font(run)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)


def create_numbering(doc: Document, ordered: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1." if ordered else "•")
    level.append(level_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "space")
    level.append(suffix)
    p_pr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def add_bullets(doc: Document, items: list[str]) -> None:
    num_id = create_numbering(doc, ordered=False)
    for item in items:
        paragraph = doc.add_paragraph()
        apply_numbering(paragraph, num_id)
        run = paragraph.add_run(item)
        set_run_font(run, 10.3)
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.first_line_indent = Cm(-0.3)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.keep_together = True


def add_numbered(doc: Document, items: list[str]) -> None:
    num_id = create_numbering(doc, ordered=True)
    for item in items:
        paragraph = doc.add_paragraph()
        apply_numbering(paragraph, num_id)
        run = paragraph.add_run(item)
        set_run_font(run, 10.3)
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.first_line_indent = Cm(-0.3)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.keep_together = True


def fill_cell(cell, text: str, bold: bool = False, color: str = INK, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(text)
    set_run_font(run, 9.6, bold=bold, color=color)


def add_metric_table(doc: Document) -> None:
    rows = [
        ("운영 데이터", "큐레이션 논문 1,651편 · 게재지 330개 · 조사 연도 2024-2026"),
        ("최신성", "최근 7일 신규 9편 · 정기 수집 1일 4회(약 6시간 간격)"),
        ("추천·트렌드", "AML 추천 741편 · 키워드 네트워크 30개 노드/130개 연결"),
        ("서비스 효율", "초기 논문 데이터 전송량 12.42MB → 2.46MB로 약 80% 절감"),
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [2200, 7240])
    set_table_borders(table)
    header = table.rows[0].cells
    fill_cell(header[0], "검증 항목", True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(header[1], "2026.09.12 기준 구현 현황", True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(header[0], ACCENT)
    set_cell_shading(header[1], ACCENT)
    set_repeat_table_header(table.rows[0])
    for label, value in rows:
        cells = table.add_row().cells
        fill_cell(cells[0], label, True, ACCENT)
        fill_cell(cells[1], value)
        set_cell_shading(cells[0], "F4F8F9")
    set_table_geometry(table, [2200, 7240])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_risk_table(doc: Document) -> None:
    rows = [
        ("AI 환각·과도한 단정", "논문 내용을 잘못 이해하거나 결과를 과장할 수 있음", "Topic·Problem·Method·Key Result·Takeaway의 고정 Q5 출력, AI/Metadata 출처 표기, DOI 원문 확인 링크, 향후 샘플 정답셋 기반 자동 평가"),
        ("분류·추천 편향", "특정 키워드·게재지에 결과가 치우칠 수 있음", "결정론적 분류 규칙과 임베딩 추천을 분리하고, AML 점수 구성(의미 80%·최신성 10%·게재지 10%)을 공개하며 수동 seed와 archive로 교정"),
        ("저작권·원문 재현", "초록 또는 PDF를 사실상 대체할 위험", "출판사 페이지 크롤링 및 PDF 저장 금지, raw abstract 비공개·비저장, 요약은 변형·압축문으로 작성하고 DOI로 출처 연결"),
        ("API 장애·호출 제한", "수집 누락 또는 자동화 지연 가능", "요청 속도·페이지 예산 제한, 실패 단위별 재시도, 기존 논문 보존, 실행 전 백업, 상태 패널과 workflow 상태 파일로 가시화"),
        ("비용·컴퓨팅 의존", "상용 LLM 비용 또는 로컬 PC 가동 의존", "정기 수집에서는 OpenAI를 호출하지 않고, 유료 작업은 수동 확인을 거침. 로컬 Ollama 중단 시 해당 AI 배치만 멈추고 공개 사이트와 정기 수집은 계속 운영"),
        ("지표 오해", "OA Rank를 공식 JCR 등급으로 오인할 수 있음", "화면에서 OpenAlex 기반 내부 venue signal임을 명시하고 JCR/IF가 아님을 툴팁과 운영정책에 고정"),
        ("비밀정보 노출", "API key·AML seed·디버그 자료 공개 위험", "GitHub Actions Secret 사용, Pages 배포 시 data/private·private 제외, 공개 JSON 스키마 점검 및 원문성 필드 제거"),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [1900, 2700, 4840])
    set_table_borders(table)
    headers = table.rows[0].cells
    for cell, text in zip(headers, ("위험요인", "발생 가능 영향", "대응방안")):
        fill_cell(cell, text, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, ACCENT)
    set_repeat_table_header(table.rows[0])
    for label, impact, response in rows:
        cells = table.add_row().cells
        fill_cell(cells[0], label, True, ACCENT)
        fill_cell(cells[1], impact)
        fill_cell(cells[2], response)
        set_cell_shading(cells[0], "F4F8F9")
    set_table_geometry(table, [1900, 2700, 4840])


def add_roadmap_table(doc: Document) -> None:
    rows = [
        ("1단계 · 교내 실증", "현재-3개월", "연구자·대학원생 대상 사용성 평가, 검색 성공률·DOI 클릭률·요약 오류 신고 수집, Q5 및 추천 임계값 보정"),
        ("2단계 · 개인화", "3-6개월", "사용자 관심 분야·seed 논문을 선택하는 개인 프로필, 주간 신규/AML 추천 이메일 알림, 저장 목록과 변경 이력 제공"),
        ("3단계 · 연구 인텔리전스", "6-12개월", "키워드 네트워크의 시계열 비교, 기관·연구실별 기술 포트폴리오, 온톨로지 기반 연관 주제 탐색, 평가 대시보드 구축"),
        ("4단계 · 확산", "12개월 이후", "타 공학 분야로 검색·분류 프로필을 모듈화하고, 대학 공동 활용 또는 연구지원 부서용 독립 배포 모델로 확장"),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [1900, 1500, 6040])
    set_table_borders(table)
    headers = table.rows[0].cells
    for cell, text in zip(headers, ("단계", "기간", "핵심 목표")):
        fill_cell(cell, text, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, ACCENT)
    set_repeat_table_header(table.rows[0])
    for phase, period, goal in rows:
        cells = table.add_row().cells
        fill_cell(cells[0], phase, True, ACCENT)
        fill_cell(cells[1], period, False, INK, WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(cells[2], goal)
        set_cell_shading(cells[0], "F4F8F9")
    set_table_geometry(table, [1900, 1500, 6040])


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(4)
        run = paragraph.add_run("AMLens · AI Manufacturing Literature Intelligence Platform  |  ")
        set_run_font(run, 8.2, color=MUTED)
        fld_char1 = OxmlElement("w:fldChar")
        fld_char1.set(qn("w:fldCharType"), "begin")
        instr_text = OxmlElement("w:instrText")
        instr_text.set(qn("xml:space"), "preserve")
        instr_text.text = "PAGE"
        fld_char2 = OxmlElement("w:fldChar")
        fld_char2.set(qn("w:fldCharType"), "end")
        run._r.append(fld_char1)
        run._r.append(instr_text)
        run._r.append(fld_char2)


def build() -> None:
    source_doc = Document(SOURCE)
    source_section = source_doc.sections[0]
    doc = Document()
    section = doc.sections[0]
    section.page_width = source_section.page_width
    section.page_height = source_section.page_height
    section.top_margin = source_section.top_margin
    section.right_margin = source_section.right_margin
    section.bottom_margin = source_section.bottom_margin
    section.left_margin = source_section.left_margin
    style_document(doc)

    # Reproduce the official title and instruction with portable Word styles.
    title = doc.add_paragraph()
    title.add_run('【“AI Native Campus” 제2회 생성형AI 경진대회 참가신청서】')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        set_run_font(run, 16, bold=True, color=INK)
    note = doc.add_paragraph()
    note.add_run('※ 현재 서식을 유지하며 아래 내용을 중심으로 10장 이내로 작성\n')
    note.add_run('(필요시 이미지 등 삽입 가능 / hwpx, docx, pdf 파일로 비교과시스템에 팀장이 업로드)')
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in note.runs:
        set_run_font(run, 9.5, color=MUTED)
    note.paragraph_format.space_after = Pt(12)

    add_section_heading(doc, "참가자 정보")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("○ 팀명: ")
    set_run_font(r, 10.8, bold=True)
    r = p.add_run("[직접 입력]")
    set_run_font(r, 10.8, bold=True, color="8A5A00")
    r._element.get_or_add_rPr().append(OxmlElement("w:highlight"))
    r._element.rPr[-1].set(qn("w:val"), "yellow")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("○ 참가자: ")
    set_run_font(r, 10.8, bold=True)
    r = p.add_run("[성명(소속, 교원/학생/직원 등) 직접 입력]")
    set_run_font(r, 10.8, bold=True, color="8A5A00")
    r._element.get_or_add_rPr().append(OxmlElement("w:highlight"))
    r._element.rPr[-1].set(qn("w:val"), "yellow")

    add_section_heading(doc, "서비스 기본정보")
    add_item_heading(doc, "서비스명")
    add_body(doc, "AMLens(에이엠렌즈) - AI Manufacturing Literature Intelligence Platform")

    add_item_heading(doc, "서비스 설명 및 핵심 기능")
    add_body(
        doc,
        "AMLens는 생산제조, 3D/4D 프린팅, 제조 로보틱스, AI Manufacturing 분야의 논문을 지속적으로 수집·분류하고, 생성형 AI 요약과 연구자 맞춤 추천, 최신 키워드 네트워크를 한 화면에서 제공하는 연구 인텔리전스 서비스다. 사용자는 매번 여러 학술 데이터베이스에서 같은 검색식을 반복하지 않고도 이번 주에 새로 들어온 논문부터 관심 주제, 게재지 신호, AML 추천 결과까지 빠르게 좁혀 원문 DOI로 이동할 수 있다.",
    )
    add_bullets(doc, [
        "자동 수집·큐레이션: Crossref 공식 API에서 2024년 이후 논문을 정기 수집하고 DOI/제목 기반 중복 제거, 관련성 점수, 분야·서브토픽 분류를 수행한다.",
        "고정형 Q5 AI 요약: 논문을 Topic, Problem, Method, Key Result, Takeaway의 다섯 질문으로 정리해 연구자가 핵심을 같은 기준으로 비교할 수 있게 한다.",
        "다층 탐색 UI: 키워드 검색, 분야·서브토픽, 연도, OA Rank, 요약 유형, 신규 여부 필터와 정렬을 결합한다. 초기 화면에는 주간 신규 논문 20편만 보여 주고 필요할 때 추가 로딩한다.",
        "AML 추천 엔진: 연구실의 핵심 seed 논문과 후보 논문의 임베딩 유사도를 비교하고 최신성·게재지 신호를 합산해 별도 AML 점수와 추천 이유를 제공한다.",
        "트렌드 키워드 네트워크: 제목·태그·Q5 요약·공개 메타데이터에서 파생한 키워드의 동시출현 관계를 움직이는 네트워크로 표현하고, 노드 클릭 시 관련 논문으로 바로 필터링한다.",
        "LLM Helper: 현재 화면의 필터와 상위 논문 정보를 근거가 포함된 프롬프트로 변환하여, 사용자가 학교 제공 GPT나 로컬 LLM에서 비교·요약·연구 질문 도출을 이어갈 수 있게 한다.",
    ])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_item_heading(doc, "현재 구현 현황")
    add_metric_table(doc)

    add_item_heading(doc, "이 서비스를 구상한 이유")
    add_body(
        doc,
        "생산제조 연구자는 논문 자체가 부족해서가 아니라, 여러 플랫폼에 흩어진 최신 논문을 지속적으로 확인하고 자신의 연구 맥락에서 우선순위를 판단하는 데 많은 시간을 쓴다. 일반 검색 서비스는 범위가 지나치게 넓고, 단순 키워드 알림은 '왜 중요한지'와 '내 연구에 얼마나 가까운지'를 설명하지 못한다. 특히 Multi-material AM, Functionally Graded AM, material switching, toolpath strategy처럼 용어가 겹치거나 빠르게 변하는 분야에서는 검색식 관리와 중복 검토가 반복 업무가 된다.",
    )
    add_body(
        doc,
        "AMLens는 이 문제를 검색 결과의 양이 아니라 '연구자가 다음에 읽을 논문을 결정하는 시간'의 문제로 정의했다. 공식 메타데이터 수집, 규칙 기반 분류, 생성형 AI Q5 요약, seed 기반 임베딩 추천을 연결해 발견-이해-선별-원문 확인의 흐름을 하나의 서비스로 묶었다. 또한 정기 수집에는 유료 LLM을 사용하지 않고, 로컬 모델과 수동 승인형 OpenAI 작업을 병행해 대학·연구실이 비용과 데이터 통제권을 유지하도록 설계했다.",
    )

    add_section_heading(doc, "서비스 상세정보", page_break=True)
    add_item_heading(doc, "서비스(시스템) 구성")
    add_numbered(doc, [
        "수집 계층 - GitHub Actions가 약 6시간 간격으로 Crossref 검색어와 선택 게재지 ISSN 질의를 실행한다. OpenAlex 일반 논문 검색은 사용하지 않으며, DOI가 있는 논문의 교신저자 누락을 교차 확인하거나 venue signal을 보강할 때만 사용한다.",
        "정제·큐레이션 계층 - DOI 우선, 제목·연도·제1저자 보조 키로 중복을 제거한다. 생산제조 맥락, 연도, 문서 유형, 게재지 신뢰 신호와 관련성 기준으로 active와 archive를 분리하되 기존 active 논문은 최신 API 응답에 없더라도 보존한다.",
        "분류·요약 계층 - Production/Manufacturing, 3D Printing, 4D Printing, Robotics for Manufacturing, AI Manufacturing의 5개 상위 분야와 MMAM, FGAM, DLP, LCE, Digital Twins, Manufacturing Automation 등 서브토픽을 부여한다. Q5 요약은 OpenAI 또는 로컬 Ollama가 생성하고, 미처리 항목은 metadata summary로 구분한다.",
        "AML 추천 계층 - 내부 탐색(score_existing)은 현재 논문 풀을, 외부 탐색(collect_and_score/full_refresh)은 venue 제한 없이 Crossref 키워드 검색 후보를 포함한다. 핵심 seed 집합의 임베딩과 후보 임베딩을 비교해 AML 점수를 계산하고 공개 임계값 이상만 추천 목록에 저장한다.",
        "프론트엔드 계층 - 빌드 서버가 필요 없는 HTML/CSS/JavaScript 정적 앱이 GitHub Pages에서 동작한다. 첫 방문은 경량 index만 받고, 사용자가 세부 내용을 열 때 Q5 요약·저자·교신저자 정보를 chunk 단위로 지연 로딩한다.",
        "운영·배포 계층 - 수집 시작/성공/실패 상태를 별도 JSON과 상태 패널에 기록하고, 데이터 변경 시 자동 커밋한 뒤 Pages artifact를 배포한다. 매 실행 전 이전 데이터셋을 압축 보관하여 장애 시 복구 가능성을 확보한다.",
    ])

    add_item_heading(doc, "사용자 경험 흐름")
    add_numbered(doc, [
        "접속 즉시 최근 7일 신규 논문 중 최대 20편을 확인한다.",
        "관심 분야·서브토픽 또는 검색어를 선택해 전체 큐레이션 데이터로 탐색 범위를 넓힌다.",
        "논문 카드에서 Q5 요약, 관련성/AML 점수, 대표 태그, 저자·교신저자, OA Rank를 비교한다.",
        "Trend Map에서 최근 강도가 높은 키워드와 연결 구조를 확인하고 노드를 클릭해 논문 목록을 교차 탐색한다.",
        "AML Recommendations를 선택해 연구실 seed와 의미적으로 가까운 논문을 별도 목록으로 검토한다.",
        "필요한 결과를 LLM Helper 프롬프트로 내보내 심층 비교 질문을 만들고, 최종 판단은 DOI의 공식 원문에서 확인한다.",
    ])

    add_item_heading(doc, "사용 모델 및 기술, 플랫폼 등 기술적 사항")
    add_bullets(doc, [
        "데이터/API: Crossref Works API를 논문 발견의 단일 출처로 사용한다. OpenAlex Works DOI lookup은 교신저자 보완, OpenAlex Sources는 OA Rank 산출에 한정한다. 출판사 웹페이지를 크롤링하지 않는다.",
        "생성 모델: 수동 OpenAI 요약은 기본 gpt-4o-mini를 사용하며 비용 확인 입력이 있어야 실행된다. 로컬 요약 에이전트는 Ollama의 qwen2.5:7b를 사용하고, 각 카드 하단에 사용 출처와 모델을 표시한다.",
        "임베딩 모델: AML 추천은 OpenAI text-embedding-3-small 또는 로컬 nomic-embed-text를 선택할 수 있다. 공급자별 벡터 차원과 점수 분포가 달라 캐시와 공개 임계값을 분리한다.",
        "AML 점수: 의미 유사도 80%, 최신성 10%, 게재지 10%의 설명 가능한 결합 점수다. 의미 유사도는 seed 평균 프로필, 가장 가까운 seed, seed 전체 평균과의 코사인 유사도를 조합한다. 키워드는 후보 발견에는 사용하지만 최종 점수에 중복 반영하지 않는다.",
        "AI 에이전트 방식: 로컬 파이프라인이 데이터 상태를 읽고 요약 대상 선별→Q5 생성→형식 검증→출처 기록→split 재생성의 도구 호출 순서를 수행한다. AML은 후보 수집→임베딩 재사용/생성→점수화→추천 이유 생성→검수 가능한 공개 JSON 작성 순으로 실행한다.",
        "웹/자동화: Python 3.11, requests, openpyxl, OpenAI SDK, 순수 JavaScript/HTML/CSS, GitHub Actions, GitHub Pages를 사용한다. 클라이언트에는 API key가 포함되지 않는다.",
        "성능: 12.42MB의 전체 active JSON 대신 2.46MB index를 먼저 받아 약 80%를 줄이고, 상세 chunk와 20편 단위 Load more를 이용해 정적 호스팅의 트래픽과 렌더링 부담을 낮춘다.",
        "추적성: summary_provider, summary_model, abstract_used_for_summary, corresponding_author_source, first_added, last_updated 등 provenance 필드로 데이터와 AI 처리 출처를 남긴다.",
    ])

    add_item_heading(doc, "생성형 AI 활용의 차별성")
    add_bullets(doc, [
        "일반 챗봇처럼 질문마다 전체 DB를 다시 읽히지 않고, 수집·분류·임베딩·요약 결과를 재사용하는 상태 기반 파이프라인이므로 반복 비용을 줄인다.",
        "Q5 출력 형식을 고정해 논문마다 문장 길이와 관점이 달라도 동일한 질문으로 비교할 수 있으며, LLM의 자유로운 서술을 연구 검토에 적합한 구조로 제한한다.",
        "추천 결과를 일반 relevance와 분리된 AML score로 저장해 '사이트 주제에 맞는 논문'과 '특정 연구실 프로필에 가까운 논문'을 구분한다.",
        "상용 모델과 로컬 모델을 동일 인터페이스로 운용하고 모델 출처를 표시하여, 품질·비용·보안 요구에 따라 교체 가능한 대학형 생성 AI 운영 모델을 제시한다.",
    ])

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_item_heading(doc, "예상되는 위험성 및 대응방안")
    add_risk_table(doc)

    add_section_heading(doc, "서비스의 기대효과", page_break=True)
    add_item_heading(doc, "예상 사용자 규모")
    add_body(
        doc,
        "초기 실증 대상은 교내 생산제조·적층제조·로보틱스·AI 연구실의 교수, 대학원생, 학부연구생과 연구지원 인력 약 100명이다. 공개 서비스와 주간 알림을 안정화한 뒤 타 대학 및 기업 R&D 연구자까지 월간 활성 사용자 1,000명 수준을 1차 확장 목표로 설정한다. 정적 웹 구조이므로 사용자별 서버 세션 없이 읽기 중심 트래픽을 수용하고, AI 연산은 배치 처리해 사용자 증가와 모델 비용이 직접 비례하지 않도록 한다.",
    )

    add_item_heading(doc, "기대효과")
    add_bullets(doc, [
        "연구 탐색 시간 단축: 새 논문 수집, 중복 제거, 분야 분류, 핵심 Q5 확인을 자동화해 연구자는 검색식 반복보다 읽을 논문 선택과 연구 설계에 집중할 수 있다.",
        "연구 품질 향상: 동일한 Q5 기준과 DOI 원문 연결로 제목만 보고 판단하는 오류를 줄이고, 방법·결과·한계를 빠르게 비교하는 습관을 지원한다.",
        "연구실 지식의 자산화: AML seed 집합과 추천 결과가 연구실의 관심 영역을 재현 가능한 데이터로 바꾸며, 신입 연구자 온보딩과 공동 연구 주제 발굴에 활용될 수 있다.",
        "최신 연구 동향의 가시화: 주간 신규량과 키워드 연결망을 함께 보여 단순 빈도보다 주제 간 결합과 새롭게 강해지는 연구 흐름을 파악할 수 있다.",
        "생성형 AI 비용 절감: 정기 수집·기본 분류는 무과금 경로로 유지하고, 로컬 GPU 요약·임베딩을 활용하며, 고품질 상용 모델은 필요한 배치만 승인 실행한다.",
        "책임 있는 학술 AI 운영: 원문·PDF를 호스팅하지 않고 모델 출처와 지표 한계를 밝히며, AI 요약을 원문 대체가 아닌 큐레이션 보조 도구로 위치시킨다.",
    ])

    add_item_heading(doc, "성과 측정 계획")
    add_bullets(doc, [
        "탐색 효율: 첫 접속부터 유효 DOI 클릭까지 걸린 시간, 검색/필터 후 원문 이동률, 주간 신규 논문 확인 완료율을 측정한다.",
        "추천 품질: 연구자 평가 기반 Precision@20, seed 대비 NDCG, 추천 숨김·저장 비율과 AML 점수 구간별 만족도를 추적한다.",
        "요약 품질: Q5 형식 통과율, 근거 없는 주장 신고율, 원문 대조 샘플의 사실 일치도와 close-paraphrase 점검 통과율을 관리한다.",
        "운영 안정성: 정기 workflow 성공률, API 오류율, 배포 지연, 초기 데이터 크기, 페이지 로딩 실패율을 상태 파일과 로그로 확인한다.",
    ])

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_item_heading(doc, "향후 확장 계획")
    add_roadmap_table(doc)
    add_body(
        doc,
        "최종적으로 AMLens를 '논문을 대신 읽어 주는 서비스'가 아니라, 연구자가 최신 문헌을 놓치지 않고 자신의 연구 방향과 연결해 판단하도록 돕는 검증 가능한 AI 연구 동료로 발전시키고자 한다. 핵심 원칙은 자동화 범위를 넓히되 원문 확인, 개인정보·저작권 보호, 모델 출처 공개와 사람의 최종 판단을 유지하는 것이다.",
    )

    add_footer(doc)
    doc.core_properties.title = "AI Native Campus 제2회 생성형AI 경진대회 참가신청서 - AMLens"
    doc.core_properties.subject = "AMLens AI Manufacturing Literature Intelligence Platform"
    doc.core_properties.author = "참가팀 직접 입력"
    doc.core_properties.last_modified_by = "참가팀 직접 입력"
    doc.core_properties.keywords = "생성형AI, AI Manufacturing, Additive Manufacturing, 논문추천, Ollama, AML"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
